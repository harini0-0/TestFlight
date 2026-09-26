"""LLM, storage, outbox, and telemetry interfaces."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol


class LlmClient(Protocol):
    def complete(self, *, model_id: str, system: str, user: str, schema_name: str) -> str: ...


class Embedder(Protocol):
    def embed(self, text: str) -> list[float]: ...


class DocumentParser(Protocol):
    def extract_text(self, payload: bytes, filename: str) -> str: ...


class S3BlobStore:
    """Contract blobs in S3. SSE-KMS is applied when GLASSWING_S3_SSE is set."""

    def __init__(self, bucket: str, endpoint_url: str | None = None) -> None:
        import boto3

        self.bucket = bucket
        self.client = boto3.client("s3", endpoint_url=endpoint_url)

    def put(self, tenant_id: str, document_id: str, filename: str, payload: bytes) -> str:
        import os

        key = f"{tenant_id}/contracts/{document_id}/{filename}"
        extra: dict = {}
        if os.environ.get("GLASSWING_S3_SSE"):
            extra["ServerSideEncryption"] = "aws:kms"
        self.client.put_object(Bucket=self.bucket, Key=key, Body=payload, **extra)
        return f"s3://{self.bucket}/{key}"

    def get(self, uri: str) -> bytes:
        bucket, key = uri.removeprefix("s3://").split("/", 1)
        body = self.client.get_object(Bucket=bucket, Key=key)["Body"].read()
        return body


@dataclass
class LocalBlobStore:
    root: Path

    def put(self, tenant_id: str, document_id: str, filename: str, payload: bytes) -> str:
        folder = self.root / tenant_id / "contracts" / document_id
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / filename
        path.write_bytes(payload)
        return f"file://{path}"

    def get(self, uri: str) -> bytes:
        path = uri.removeprefix("file://")
        return Path(path).read_bytes()


def looks_like_document(payload: bytes, filename: str) -> bool:
    name = filename.lower()
    if name.endswith(".txt") or name.endswith(".csv"):
        return True
    if payload.startswith(b"%PDF") or payload.startswith(b"PK"):
        return True
    if name.endswith(".pdf") or name.endswith(".docx") or name.endswith(".xlsx") or name.endswith(".xlsm"):
        return payload.startswith(b"%PDF") or payload.startswith(b"PK")
    return False


def _is_xlsx(payload: bytes, filename: str) -> bool:
    name = filename.lower()
    return name.endswith((".xlsx", ".xlsm")) and payload.startswith(b"PK")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\t", " ").replace("\n", " ").strip()


def _xlsx_text(payload: bytes) -> str:
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for sheet in book.worksheets:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell_text(value) for value in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip())
            if lines:
                sections.append(f"Sheet: {sheet.title}\n" + "\n".join(lines))
    finally:
        book.close()
    return "\n\n".join(sections)


_PDF_DUMP = re.compile(r"(^%PDF-)|(\bendobj\b)|(\bendstream\b)|(\bstartxref\b)", re.MULTILINE)
_PDF_NOISE_LINE = re.compile(
    r"^\s*(%PDF-|xref\b|trailer\b|startxref\b|endobj\b|endstream\b|\d+\s+\d+\s+obj\b|/Type\b|/Filter\b)"
)


def _is_visible_char(ch: str) -> bool:
    code = ord(ch)
    if ch in "\t ":
        return True
    if code < 32 or code == 127 or 0x80 <= code <= 0x9F:
        return False
    if 0xE000 <= code <= 0xF8FF or 0xF0000 <= code <= 0x10FFFD:
        return False
    if code in {0xFFFD, 0xFFFE, 0xFFFF, 0x200B, 0x200C, 0x200D, 0x2060, 0xFEFF}:
        return False
    return True


def _is_pdf_dump(text: str) -> bool:
    head = text.lstrip()[:400]
    if head.startswith("%PDF-"):
        return True
    return "endobj" in head and "/Type" in head


def visible_text(text: str) -> str:
    """Words and numbers a person can read. Drops file syntax and hidden glyphs."""
    if not text or _is_pdf_dump(text):
        return ""
    lines: list[str] = []
    blank = False
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if _PDF_NOISE_LINE.match(raw) or (_PDF_DUMP.search(raw) and "<<" in raw):
            continue
        kept = [" " if ch == "\u00a0" else ch for ch in raw if _is_visible_char(ch)]
        line = "".join(kept).rstrip(" \t")
        if not line.strip():
            hidden_only = bool(raw.strip()) and not any(ch not in "\t " and _is_visible_char(ch) for ch in raw)
            if hidden_only:
                continue
            if lines and not blank:
                lines.append("")
                blank = True
            continue
        blank = False
        lines.append(line)
    return "\n".join(lines).strip()


def _pdf_text(payload: bytes) -> str:
    import io

    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(payload))
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception:
        return ""
    if text.lstrip().startswith("%PDF"):
        return ""
    return text


class LocalPdfParser:
    def extract_text(self, payload: bytes, filename: str) -> str:
        if _is_xlsx(payload, filename):
            try:
                raw = _xlsx_text(payload)
            except Exception:
                raw = ""
        elif payload.startswith(b"%PDF"):
            raw = _pdf_text(payload)
        else:
            raw = payload.decode("utf-8", errors="ignore")
        return visible_text(raw)


@dataclass
class OutboxMessage:
    event_id: str
    event_type: str
    payload: dict
    published: bool = False


@dataclass
class Outbox:
    messages: list[OutboxMessage] = field(default_factory=list)

    def add(self, event_id: str, event_type: str, payload: dict) -> None:
        self.messages.append(OutboxMessage(event_id, event_type, payload))

    def relay(self, publisher) -> int:
        count = 0
        for message in self.messages:
            if message.published:
                continue
            publisher(message.event_type, message.payload, message.event_id)
            message.published = True
            count += 1
        return count


class MemoryPublisher:
    def __init__(self) -> None:
        self.events: list[tuple[str, dict, str]] = []

    def __call__(self, event_type: str, payload: dict, event_id: str) -> None:
        self.events.append((event_type, payload, event_id))


def hash_api_key(raw: str) -> str:
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def configure_telemetry(service_name: str) -> None:
    """OpenTelemetry is a no-op unless OTEL_EXPORTER_OTLP_ENDPOINT is set."""
    import os

    endpoint = os.environ.get("OTEL_EXPORTER_OTLP_ENDPOINT")
    if not endpoint:
        return
    try:
        from opentelemetry import trace
        from opentelemetry.exporter.otlp.proto.http.trace_exporter import OTLPSpanExporter
        from opentelemetry.sdk.resources import Resource
        from opentelemetry.sdk.trace import TracerProvider
        from opentelemetry.sdk.trace.export import BatchSpanProcessor
    except ImportError:
        return
    provider = TracerProvider(resource=Resource.create({"service.name": service_name}))
    provider.add_span_processor(BatchSpanProcessor(OTLPSpanExporter(endpoint=endpoint)))
    trace.set_tracer_provider(provider)


def langfuse_trace(name: str, prompt_hash: str, model_id: str, latency_ms: int) -> None:
    import os

    host = os.environ.get("LANGFUSE_HOST")
    if not host:
        return
    payload = {
        "name": name,
        "metadata": {"prompt_hash": prompt_hash, "model_id": model_id, "latency_ms": latency_ms},
    }
    try:
        import urllib.request

        request = urllib.request.Request(
            f"{host.rstrip('/')}/api/public/ingestion",
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
        )
        urllib.request.urlopen(request, timeout=2)  # noqa: S310
    except Exception:
        return
